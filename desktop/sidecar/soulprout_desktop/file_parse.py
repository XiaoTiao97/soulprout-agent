"""本地文档解析：本地会话读 pdf / docx / pptx / xlsx 时在用户电脑上解析，不再把文件传到云端。

和云端 ``AsyncFileProcess`` 一样，只抽取文字内容。依赖都是可选的（pdfminer.six、python-docx、python-pptx、
openpyxl），缺哪个就返回一条明确的提示，不会让整个工具崩掉。旧版 ``.doc`` 本地不支持。
"""

from __future__ import annotations

import asyncio
from typing import Callable, Dict

PARSEABLE_SUFFIXES = (".pdf", ".docx", ".doc", ".pptx", ".xlsx")


def _missing(pkg: str, suffix: str) -> str:
    return f"错误：本机缺少解析 {suffix} 所需的组件（{pkg}），无法读取该文件。"


def _parse_pdf(path: str) -> str:
    try:
        from pdfminer.high_level import extract_text
    except ImportError:
        return _missing("pdfminer.six", ".pdf")
    return extract_text(path) or ""


def _parse_docx(path: str) -> str:
    try:
        from docx import Document
    except ImportError:
        return _missing("python-docx", ".docx")
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_pptx(path: str) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return _missing("python-pptx", ".pptx")
    prs = Presentation(path)
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        slides.append(f"--- 第 {index} 页 ---\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _parse_xlsx(path: str) -> str:
    try:
        import openpyxl
    except ImportError:
        return _missing("openpyxl", ".xlsx")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    rows.append(" | ".join("" if cell is None else str(cell) for cell in row))
            sheets.append(f"--- 工作表 {sheet.title} ---\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(sheets)


_PARSERS: Dict[str, Callable[[str], str]] = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".pptx": _parse_pptx,
    ".xlsx": _parse_xlsx,
}


async def parse_document(path: str, display_name: str) -> str:
    suffix = "." + display_name.rsplit(".", 1)[-1].lower() if "." in display_name else ""
    if suffix == ".doc":
        return "错误：本地暂不支持解析旧版 .doc 文件，请先另存为 .docx。"
    parser = _PARSERS.get(suffix)
    if parser is None:
        return f"错误：不支持解析的文件类型：{suffix or display_name}"
    return await asyncio.to_thread(parser, path)
